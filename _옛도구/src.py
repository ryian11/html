# -*- coding: utf-8 -*-
"""Lesson 8 원고 추출: 사운드 대본(문장↔mp3) + 스토리보드 Dictation(빈칸)"""
import openpyxl, re

SND='/mnt/user-data/uploads/2026_cj_midd3_eng/02_사운드/중3(소영순) 전자저작물 녹음 대본_최종_20260806 아이스캔디 전달.xlsx'
SB ='/mnt/user-data/uploads/2026_cj_midd3_eng/01_스토리보드/전자저작물 추가 원고_20260716_아이스캔디 전달/3학년 전자저작물_지시문_딕테이션_미니 단어장_구문 해설_Lesson %d.xlsx'

def norm(s):
    s = s.replace('\u2019','’').replace('\u201c','“').replace('\u201d','”')
    return re.sub(r'\s+',' ', s).strip()

def sound_sentences(lesson, page):
    """[(mp3파일명, 문장)] — 화자표시(W:/M:) 제거"""
    wb=openpyxl.load_workbook(SND, data_only=True)
    ws=wb['%d과'%lesson]
    out=[]
    pat=re.compile(r'^3-%s-Read-(\d+)$'%page)
    started=False
    for r in ws.iter_rows(min_row=1, values_only=True):
        fn=str(r[2] or '').strip()
        txt=str(r[1] or '')
        m=pat.match(fn)
        if m:
            started=True
            out.append(['3_%s_read_%s.mp3'%(page, m.group(1)), norm(_spk(txt))])
        elif started and not fn and txt.strip():
            out[-1][1] = norm(out[-1][1] + ' ' + _spk(txt))   # 파일명 없는 줄 = 앞 음원에 이어붙임
        elif started and fn:
            break
    return [tuple(x) for x in out]

def _spk(s):
    return re.sub(r'^\s*[A-Z][a-z]*\s*:\s*','',s)

def runs(v):
    if v is None: return []
    if isinstance(v,str): return [(v,False)]
    out=[]
    for b in v:
        if isinstance(b,str): out.append((b,False))
        else:
            c=getattr(getattr(getattr(b,'font',None),'color',None),'rgb',None)
            out.append((b.text, str(c)=='FFFF0000'))
    return out

def dictation(lesson):
    """{섹션: {난이도: [문단]}}  문단 = [(글자, 빨강여부)] 런 목록"""
    wb=openpyxl.load_workbook(SB%lesson, data_only=True, rich_text=True)
    ws=wb['Dictation']
    data={}
    for row in range(2, ws.max_row+1):
        sec=ws.cell(row=row,column=1).value
        lv =ws.cell(row=row,column=2).value
        if not sec or not lv: continue
        rr=runs(ws.cell(row=row,column=3).value)
        paras=_split_paras(rr)
        data.setdefault(str(sec).strip(), {})[str(lv).strip()] = paras
    return data

def to_marked(para_runs):
    """런 목록 → '||단어|| 포함 평문'.
    빨강 런 하나 = 정답 버튼 하나. 런 안의 단어는 각각 별도 빈칸(‖로 런 경계 표시)."""
    out=[]
    for text, red in para_runs:
        text = text.replace('\n',' ')
        if red:
            if out and out[-1][1]:
                out[-1] = (out[-1][0] + text, True)
            else:
                out.append((text, True))
        else:
            out.append((text, False))
    s=''
    for text, red in out:
        if red:
            t=text.strip()
            lead=' ' if text[:1]==' ' else ''
            tail=' ' if text[-1:]==' ' else ''
            s += lead + '\u2016' + ' '.join('||%s||'%w for w in t.split()) + '\u2016' + tail
        else:
            s += text
    return re.sub(r'[ \t]+',' ', s).strip()

def strip_marks(s):
    return norm(re.sub(r'\|\|([^|]*)\|\|', r'\1', s.replace('\u2016','')))

def split_groups(marked):
    """'‖...‖' 로 감싼 구간 = 정답 버튼 하나. 반환: (en문자열, [그룹별 빈칸수])"""
    sizes=[]
    def rep(m):
        sizes.append(len(re.findall(r'\|\|[^|]*\|\|', m.group(1))))
        return m.group(1)
    en=re.sub(r'\u2016(.*?)\u2016', rep, marked)
    return re.sub(r'[ \t]+',' ',en).strip(), sizes

def item_idx(sizes, start):
    """itemIdx 문자열과 다음 시작 번호"""
    groups=[]; n=start
    for sz in sizes:
        groups.append(','.join(str(n+k) for k in range(sz))); n+=sz
    return '|'.join(groups), n

def align(marked_paras, sentences):
    """문단별 마킹 텍스트를 mp3 문장 순서에 맞춰 자른다.
       반환: [[(mp3, 마킹문장)] , ...]  문단 단위"""
    res, si = [], 0
    for mp in marked_paras:
        rest, group = mp, []
        while si < len(sentences) and rest.strip():
            mp3, sent = sentences[si]
            take, rest2 = _cut(rest, sent)
            if take is None: break
            group.append((mp3, take))
            rest = rest2; si += 1
        res.append(group)
    return res, si

def _cut(marked, sentence):
    """marked 앞부분에서 sentence 하나만큼 잘라낸다(마킹 유지)."""
    target = norm(sentence)
    plain_i = 0; out = []; i = 0
    tl = len(target)
    while i < len(marked):
        if marked.startswith('||', i):
            j = marked.index('||', i+2) + 2
            chunk = marked[i:j]
            inner = chunk[2:-2]
            out.append(chunk)
            plain_i += len(norm(inner)) if plain_i==0 else len(inner)
            i = j
        else:
            out.append(marked[i]); i += 1
            plain_i += 1
        if norm(strip_marks(''.join(out))) == target:
            return ''.join(out).strip(), marked[i:]
        if len(norm(strip_marks(''.join(out)))) > tl + 2:
            return None, marked
    return None, marked


def _split_paras(rr):
    """문단 구분 = 줄 시작에 공백이 있으면 새 문단, 없으면 앞 줄에 이어붙임.
    (한글/엑셀 원고에서 문단 들여쓰기로 구분하고, 칸 폭 때문에 문장 중간에서 줄이 끊기는 경우 대응)"""
    paras, cur, line_start = [], [], True
    for text, red in rr:
        pieces = text.split('\n')
        for k, piece in enumerate(pieces):
            if k:
                line_start = True
            if line_start and piece[:1] in (' ', '\t'):
                if cur: paras.append(cur)
                cur = []
            elif k and piece:
                if cur: cur.append((' ', False))
            if piece:
                cur.append((piece, red))
                line_start = False
    if cur: paras.append(cur)
    return [p for p in paras if ''.join(t for t,_ in p).strip()]
